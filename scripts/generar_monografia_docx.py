#!/usr/bin/env python3
"""Ensambla la monografía final (.docx) sobre la plantilla oficial `templates/apa7ma.docx`.

Es GENÉRICO: no contiene ningún tema, título ni fuente incrustados. Toma el cuerpo
en Markdown y lo vuelca con los estilos de la plantilla (Heading 1/2/3, que ya traen
numeración automática multinivel — por eso los títulos se escriben SIN números).

Uso:
  python scripts/generar_monografia_docx.py --input output/monografia.md \
      [--out output/monografia.docx] [--template templates/apa7ma.docx] \
      [--references output/referencias.json] [--caratula output/caratula.json]

Reglas que aplica:
  - `#` → Heading 1 · `##` → Heading 2 · `###` → Heading 3 · `####` → Heading 4
  - párrafos con sangría APA (0,5"), doble espacio, justificados
  - lista con viñetas · tablas con bordes APA 7
  - Sección "Resumen": sin sangría de primera línea
  - Sección "Referencias"/"Bibliografía": sangría francesa 1,27 cm
  - `*cursiva*` y `**negrita**` en línea
Requiere: python-docx (pip install python-docx)
"""
import argparse
import json
import re
from pathlib import Path

import docx
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt

BASE_DIR = Path(__file__).resolve().parent.parent

SECCION_SIN_SANGRIA = {"resumen"}
SECCION_FRANCESA = {"referencias", "referencias bibliograficas", "bibliografia"}
HEADING_STYLE = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3", 4: "Heading 4", 5: "Heading 5", 6: "Heading 6"}


# ---------------------------------------------------------------- utilidades Word
def _set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tb = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ("top", "left", "bottom", "right"):
        d = kwargs.get(edge)
        if d:
            tag = f"w:{edge}"
            tb.append(parse_xml(
                f'<{tag} {nsdecls("w")} w:val="{d.get("val","single")}" '
                f'w:sz="{d.get("sz",4)}" w:space="0" w:color="{d.get("color","auto")}"/>'))
    tcPr.append(tb)


def _apa_borders(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    n = len(table.rows)
    for r_idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if r_idx == 0:
                _set_cell_border(cell, top={"sz": 8}, bottom={"sz": 6}, left={"val": "nil"}, right={"val": "nil"})
            elif r_idx == n - 1:
                _set_cell_border(cell, bottom={"sz": 8}, left={"val": "nil"}, right={"val": "nil"}, top={"val": "nil"})
            else:
                _set_cell_border(cell, top={"val": "nil"}, bottom={"val": "nil"}, left={"val": "nil"}, right={"val": "nil"})


def _add_runs(p, text, size=12):
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|__[^_]+__|_[^_]+_)", text)
    for tk in parts:
        if not tk:
            continue
        t, bold, italic = tk, False, False
        if tk.startswith("**") and tk.endswith("**") and len(tk) > 4 or tk.startswith("__") and tk.endswith("__") and len(tk) > 4:
            t, bold = tk[2:-2], True
        elif tk.startswith("*") and tk.endswith("*") and len(tk) > 2 or tk.startswith("_") and tk.endswith("_") and len(tk) > 2:
            t, italic = tk[1:-1], True
        run = p.add_run(t)
        run.bold = bold
        run.italic = italic
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)


def _body_paragraph(doc, text, *, indent=True, hanging=False, size=12):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    pf = p.paragraph_format
    pf.line_spacing = 2.0
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if hanging:
        pf.left_indent = Inches(0.5)
        pf.first_line_indent = Inches(-0.5)
    elif indent:
        pf.first_line_indent = Inches(0.5)
    _add_runs(p, text, size)
    return p


def _bullet(doc, text, size=12):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    pf = p.paragraph_format
    pf.left_indent = Inches(0.5)
    pf.first_line_indent = Inches(-0.25)
    pf.line_spacing = 2.0
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    rb = p.add_run("• ")
    rb.font.name = "Times New Roman"
    rb.font.size = Pt(size)
    _add_runs(p, text, size)
    return p


# ---------------------------------------------------------------- parser Markdown
def parse_md(text):
    lines = text.splitlines()
    blocks, i = [], 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s.startswith("<!--"):
            while i < len(lines) and "-->" not in lines[i]:
                i += 1
            i += 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            blocks.append((f"h{len(m.group(1))}", m.group(2).strip()))
            i += 1
            continue
        if s.startswith("|") and s.count("|") >= 2:
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        if s.startswith(("- ", "* ")):
            blocks.append(("bullet", s[2:].strip()))
            i += 1
            continue
        if s.startswith("> "):
            blocks.append(("nota", s[2:].strip()))
            i += 1
            continue
        # párrafo: unir líneas consecutivas sin marca (por si el texto viene envuelto)
        buf = [s]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or re.match(r"^(#{1,6}\s|[-*>]\s|\|)", nxt):
                break
            buf.append(nxt)
            i += 1
        blocks.append(("p", " ".join(buf)))
    return blocks


def _norm(t):
    return t.strip().lower().rstrip(".:").replace("í", "i").replace("á", "a").replace("é", "e").replace("ó", "o")


# ---------------------------------------------------------------- carátula
def _fill_caratula(doc, c):
    """Completa los 10 elementos de la carátula institucional (python-docx doc.paragraphs[0..9])."""
    paras = doc.paragraphs
    campos = [
        (0, c.get("universidad", "UNIVERSIDAD AUTÓNOMA GABRIEL RENÉ MORENO")),
        (1, c.get("facultad", "FACULTAD INTEGRAL DEL NORTE")),
        (2, c.get("escuela", "ESCUELA DE INGENIERÍA Y POSTGRADO (EIP)")),
        (3, c.get("logo", "[ LOGO INSTITUCIONAL UAGRM ]")),
        (4, c.get("programa", "PROGRAMA / DIPLOMADO")),
        (5, f"MENCIÓN: {c['mencion']}" if c.get("mencion") else ""),
        (6, c.get("titulo", "TÍTULO DE LA MONOGRAFÍA")),
        (7, c.get("grado", "Monografía de Compilación presentada para optar al Certificado / Diploma correspondiente.")),
        (8, f"Postgraduante: {c['autor']}" if c.get("autor") else "Postgraduante: [Nombre completo]"),
        (9, f"Docente Guía / Tutor: {c['tutor']}" if c.get("tutor") else "Docente Guía / Tutor: [Nombre completo]"),
    ]
    for idx, val in campos:
        if idx < len(paras):
            for run in list(paras[idx].runs):
                run.text = ""
            if paras[idx].runs:
                paras[idx].runs[0].text = val
            else:
                paras[idx].add_run(val)
    # Lugar y año (elemento 10) en la primera línea vacía de la carátula.
    if c.get("lugar") or c.get("anio"):
        linea = f"{c.get('lugar', 'Santa Cruz de la Sierra, Bolivia')}, {c.get('anio', '')}".strip().strip(",")
        for idx in (10, 11):
            if idx < len(paras) and not paras[idx].text.strip():
                if paras[idx].runs:
                    paras[idx].runs[0].text = linea
                else:
                    paras[idx].add_run(linea)
                break


# ---------------------------------------------------------------- construcción
def build(args):
    tmpl = Path(args.template)
    if not tmpl.exists():
        raise SystemExit(f"✖ No se encontró la plantilla: {tmpl}")
    doc = docx.Document(str(tmpl))

    if args.caratula and Path(args.caratula).exists():
        _fill_caratula(doc, json.loads(Path(args.caratula).read_text(encoding="utf-8")))

    # Eliminar el contenido de muestra posterior a la carátula (conserva carátula + TOC/sections)
    for p in list(doc.paragraphs[12:]):
        p._p.getparent().remove(p._p)

    body = Path(args.input).read_text(encoding="utf-8")
    bloques = parse_md(body)

    if not any(k == "h1" and _norm(v) in SECCION_FRANCESA for k, v in bloques) and args.references and Path(args.references).exists():
        refs = json.loads(Path(args.references).read_text(encoding="utf-8"))
        if isinstance(refs, dict):
            refs = refs.get("fuentes", [])
        if refs:
            bloques.append(("h1", "Referencias"))
            for f in sorted(refs, key=lambda x: (x.get("formato_apa7") or "").lower()):
                bloques.append(("p", f.get("formato_apa7") or ""))

    seccion = ""
    for kind, val in bloques:
        if kind.startswith("h"):
            lvl = int(kind[1])
            h = doc.add_paragraph(val, style=HEADING_STYLE.get(lvl, "Heading 3"))
            for run in h.runs:
                run.font.name = "Times New Roman"
            seccion = _norm(val)
            continue
        if kind == "bullet":
            _bullet(doc, val)
        elif kind == "nota":
            _body_paragraph(doc, val, indent=False, size=10)
        elif kind == "table" and val:
            _render_table(doc, val)
        else:
            if "palabras clave" in val.lower():
                _body_paragraph(doc, val, indent=False, hanging=True)
            elif seccion in SECCION_FRANCESA:
                _body_paragraph(doc, val, hanging=True)
            elif seccion in SECCION_SIN_SANGRIA:
                _body_paragraph(doc, val, indent=False)
            else:
                _body_paragraph(doc, val)

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"✔ Documento generado en: {out}")


def _render_table(doc, rows):
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx in range(ncols):
            cell = t.cell(r_idx, c_idx)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            txt = row[c_idx] if c_idx < len(row) else ""
            run = p.add_run(txt)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            if r_idx == 0:
                run.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _apa_borders(t)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", required=True, help="Markdown del cuerpo de la monografía")
    ap.add_argument("--out", default=str(BASE_DIR / "output" / "monografia.docx"))
    ap.add_argument("--template", default=str(BASE_DIR / "templates" / "apa7ma.docx"))
    ap.add_argument("--references", help="output/referencias.json (si el .md no trae Referencias)")
    ap.add_argument("--caratula", help="output/caratula.json con los datos de portada")
    build(ap.parse_args())


if __name__ == "__main__":
    main()
